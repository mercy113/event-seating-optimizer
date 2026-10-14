from io import BytesIO

from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_file
)

from openpyxl import (
    Workbook,
    load_workbook
)

from models import Guest

from optimizer import (
    SeatingOptimizer
)

from utils import (
    MAX_HALL_CAPACITY,
    MAX_BALCONY_CAPACITY,
    MAX_TOTAL_CAPACITY,
    normalize_section,
    normalize_priority,
    normalize_venue,
    capacity_status
)


app = Flask(__name__)

app.config[
    "MAX_CONTENT_LENGTH"
] = 5 * 1024 * 1024


# =====================================================
# PAGES
# =====================================================

@app.route("/")
def index():

    return render_template(
        "index.html"
    )


@app.route("/setup")
def setup():

    return render_template(
        "setup.html"
    )


@app.route("/guests")
def guests_page():

    return render_template(
        "guests.html"
    )


@app.route("/result")
def result_page():

    return render_template(
        "result.html"
    )


# =====================================================
# EXCEL UPLOAD
# =====================================================

@app.post("/upload-excel")
def upload_excel():

    file = request.files.get(
        "file"
    )

    if not file:

        return jsonify(
            success=False,
            message=
            "Please select an Excel file."
        ), 400

    if not file.filename:

        return jsonify(
            success=False,
            message=
            "Please select an Excel file."
        ), 400

    if not file.filename.lower().endswith(
        (
            ".xlsx",
            ".xlsm"
        )
    ):

        return jsonify(
            success=False,
            message=
            "Only Excel files are supported."
        ), 400

    try:

        workbook = load_workbook(
            filename=BytesIO(
                file.read()
            ),
            read_only=True,
            data_only=True
        )

        worksheet = (
            workbook.active
        )

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:

            return jsonify(
                success=False,
                message=
                "Excel file is empty."
            ), 400

        headers = [
            str(value or "")
            .strip()
            .lower()
            for value in rows[0]
        ]

        aliases = {

            "name": {
                "name",
                "guest",
                "guest name",
                "full name"
            },

            "group": {
                "group",
                "guest group"
            },

            "priority": {
                "priority",
                "type"
            },

            "preferred_section": {
                "preferred section",
                "section",
                "preferred",
                "venue"
            }
        }

        def find_column(
            field
        ):

            for index, header in enumerate(
                headers
            ):

                if header in aliases[field]:

                    return index

            return None

        name_column = find_column(
            "name"
        )

        group_column = find_column(
            "group"
        )

        priority_column = find_column(
            "priority"
        )

        section_column = find_column(
            "preferred_section"
        )

        if name_column is None:

            return jsonify(
                success=False,
                message=
                "Excel must contain a Name column."
            ), 400

        guests = []

        existing_names = set()

        for row in rows[1:]:

            if (
                name_column
                >= len(row)
            ):
                continue

            name = str(
                row[name_column]
                or ""
            ).strip()

            if not name:
                continue

            key = name.lower()

            if key in existing_names:
                continue

            existing_names.add(
                key
            )

            if (
                group_column is not None
                and
                group_column < len(row)
            ):

                group = str(
                    row[group_column]
                    or "General"
                ).strip()

            else:

                group = "General"

            if (
                priority_column is not None
                and
                priority_column < len(row)
            ):

                priority = normalize_priority(
                    row[priority_column]
                )

            else:

                priority = "Standard"

            if (
                section_column is not None
                and
                section_column < len(row)
            ):

                section = normalize_section(
                    row[section_column]
                )

            else:

                section = "Any"

            guests.append({

                "name": name,

                "group":
                    group or "General",

                "priority":
                    priority,

                "preferred_section":
                    section
            })

            if (
                len(guests)
                >= MAX_TOTAL_CAPACITY
            ):
                break

        workbook.close()

        return jsonify(

            success=True,

            count=len(
                guests
            ),

            guests=guests,

            message=
            f"{len(guests)} guest(s) imported successfully."
        )

    except Exception as error:

        return jsonify(

            success=False,

            message=
            f"Could not read Excel file: {error}"
        ), 400


# =====================================================
# EXCEL TEMPLATE
# =====================================================

@app.get("/download-template")
def download_template():

    workbook = Workbook()

    worksheet = (
        workbook.active
    )

    worksheet.title = "Guests"

    worksheet.append([
        "Name",
        "Group",
        "Priority",
        "Preferred Section"
    ])

    sample_data = [

        [
            "Mercypriya",
            "Family",
            "VIP",
            "Hall"
        ],

        [
            "Priyanjali",
            "Family",
            "Standard",
            "Hall"
        ],

        [
            "Rahul",
            "Friends",
            "Standard",
            "Balcony"
        ],

        [
            "Anjali",
            "Friends",
            "VIP",
            "Balcony"
        ]

    ]

    for row in sample_data:

        worksheet.append(
            row
        )

    for column in worksheet.columns:

        maximum = 0

        for cell in column:

            length = len(
                str(
                    cell.value
                    or ""
                )
            )

            maximum = max(
                maximum,
                length
            )

        worksheet.column_dimensions[
            column[0].column_letter
        ].width = maximum + 4

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return send_file(

        output,

        as_attachment=True,

        download_name=
        "SeatSmart_Guest_Template.xlsx",

        mimetype=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# =====================================================
# OPTIMIZE
# =====================================================

@app.post("/optimize")
def optimize():

    data = (
        request.get_json(
            silent=True
        )
        or {}
    )

    venue_type = data.get(
        "venue_type",
        "both"
    )

    hall_capacity, balcony_capacity = (
        normalize_venue(

            venue_type,

            data.get(
                "hall_capacity",
                300
            ),

            data.get(
                "balcony_capacity",
                200
            )
        )
    )

    raw_guests = data.get(
        "guests",
        []
    )

    if not raw_guests:

        return jsonify(

            success=False,

            message=
            "Please add at least one guest."
        ), 400

    guests = []

    existing = set()

    for item in raw_guests:

        name = str(
            item.get(
                "name",
                ""
            )
        ).strip()

        if not name:
            continue

        key = name.lower()

        if key in existing:
            continue

        existing.add(
            key
        )

        guests.append(

            Guest(

                name=name,

                group=
                    str(
                        item.get(
                            "group",
                            "General"
                        )
                    ).strip()
                    or
                    "General",

                priority=
                    normalize_priority(
                        item.get(
                            "priority",
                            "Standard"
                        )
                    ),

                preferred_section=
                    normalize_section(
                        item.get(
                            "preferred_section",
                            "Any"
                        )
                    )
            )
        )

    status = capacity_status(

        len(guests),

        hall_capacity,

        balcony_capacity
    )

    if (
        len(guests)
        >
        hall_capacity +
        balcony_capacity
    ):

        return jsonify(

            success=False,

            status=
                status["status"],

            label=
                status["label"],

            over_capacity=
                status["over_capacity"],

            message=
                "HOUSE FULL: "
                +
                str(len(guests))
                +
                " guests for "
                +
                str(
                    hall_capacity +
                    balcony_capacity
                )
                +
                " seats. Over capacity by "
                +
                str(
                    status["over_capacity"]
                )
                +
                "."
        ), 409

    together = []

    for pair in data.get(
        "together",
        []
    ):

        if (
            isinstance(pair, list)
            and
            len(pair) >= 2
        ):

            together.append(
                (
                    str(
                        pair[0]
                    ).strip(),

                    str(
                        pair[1]
                    ).strip()
                )
            )

    apart = []

    for pair in data.get(
        "apart",
        []
    ):

        if (
            isinstance(pair, list)
            and
            len(pair) >= 2
        ):

            apart.append(
                (
                    str(
                        pair[0]
                    ).strip(),

                    str(
                        pair[1]
                    ).strip()
                )
            )

    optimizer = SeatingOptimizer(

        guests,

        together=together,

        apart=apart,

        hall_capacity=
            hall_capacity,

        balcony_capacity=
            balcony_capacity
    )

    result = optimizer.solve()

    if result is None:

        return jsonify(

            success=False,

            message=
            "No valid seating arrangement was found. "
            "Please check group, Together, Apart and "
            "section constraints."
        ), 422

    return jsonify(

        success=True,

        result=result
    )


# =====================================================
# DOWNLOAD FINAL EXCEL
# =====================================================

@app.post("/download-excel")
def download_excel():

    data = (
        request.get_json(
            silent=True
        )
        or {}
    )

    result = data.get(
        "result",
        {}
    )

    guests = data.get(
        "guests",
        []
    )

    hall_names = result.get(
        "hall",
        result.get(
            "normal",
            []
        )
    )

    balcony_names = result.get(
        "balcony",
        []
    )

    guest_map = {

        str(
            guest.get(
                "name"
            )
        ): guest

        for guest in guests

    }

    workbook = Workbook()

    worksheet = (
        workbook.active
    )

    worksheet.title = (
        "Seating Plan"
    )

    worksheet.append([

        "S.No",

        "Guest Name",

        "Group",

        "Priority",

        "Preferred Section",

        "Assigned Section",

        "Seat Number"

    ])

    serial = 1

    for seat, name in enumerate(
        hall_names,
        start=1
    ):

        guest = guest_map.get(
            name,
            {}
        )

        worksheet.append([

            serial,

            name,

            guest.get(
                "group",
                "General"
            ),

            guest.get(
                "priority",
                "Standard"
            ),

            guest.get(
                "preferred_section",
                "Any"
            ),

            "Hall",

            f"H-{seat:03d}"

        ])

        serial += 1

    for seat, name in enumerate(
        balcony_names,
        start=1
    ):

        guest = guest_map.get(
            name,
            {}
        )

        worksheet.append([

            serial,

            name,

            guest.get(
                "group",
                "General"
            ),

            guest.get(
                "priority",
                "Standard"
            ),

            guest.get(
                "preferred_section",
                "Any"
            ),

            "Balcony",

            f"B-{seat:03d}"

        ])

        serial += 1

    summary = workbook.create_sheet(
        "Summary"
    )

    summary_data = [

        [
            "SeatSmart Event Seating Optimizer",
            ""
        ],

        [
            "Hall Capacity",
            result.get(
                "hall_capacity",
                0
            )
        ],

        [
            "Hall Used",
            len(hall_names)
        ],

        [
            "Balcony Capacity",
            result.get(
                "balcony_capacity",
                0
            )
        ],

        [
            "Balcony Used",
            len(balcony_names)
        ],

        [
            "Total Guests",
            result.get(
                "total_guests",
                len(
                    guests
                )
            )
        ],

        [
            "Total Capacity",
            result.get(
                "total_capacity",
                0
            )
        ],

        [
            "Remaining Seats",
            result.get(
                "remaining_seats",
                0
            )
        ],

        [
            "Search Nodes",
            result.get(
                "nodes",
                0
            )
        ],

        [
            "Backtracks",
            result.get(
                "backtracks",
                0
            )
        ],

        [
            "Execution Time (ms)",
            result.get(
                "elapsed_ms",
                0
            )
        ],

        [
            "Efficiency (%)",
            result.get(
                "efficiency",
                0
            )
        ]

    ]

    for row in summary_data:

        summary.append(
            row
        )

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            maximum = 0

            for cell in column:

                maximum = max(
                    maximum,
                    len(
                        str(
                            cell.value
                            or ""
                        )
                    )
                )

            sheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                maximum + 3,
                45
            )

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return send_file(

        output,

        as_attachment=True,

        download_name=
        "SeatSmart_Seating_Result.xlsx",

        mimetype=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# =====================================================
# RUN
# =====================================================

if __name__ == "__main__":

    app.run(

        host="127.0.0.1",

        port=5000,

        debug=True
    )